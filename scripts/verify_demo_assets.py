import os
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from PIL import Image
import PyPDF2

def verify_xlsx():
    wb = load_workbook('data/demo_assets/sensor_readings.xlsx')
    assert 'Readings' in wb.sheetnames
    assert 'Spec_Limits' in wb.sheetnames
    
    ws = wb['Readings']
    assert ws.max_row == 402 # 400 data rows + 2 header rows
    
    # Check merged header exists
    merged = False
    for merge in ws.merged_cells.ranges:
        if str(merge) == 'A1:C1':
            merged = True
    assert merged, "Merged header A1:C1 not found"
    
    # Check out of spec and text anomaly
    text_count = 0
    out_of_spec_count = 0
    
    for row in range(3, ws.max_row + 1):
        p = ws.cell(row=row, column=2).value
        f = ws.cell(row=row, column=3).value
        
        if type(f) == str:
            text_count += 1
            
        if type(p) in (int, float) and (p < 80 or p > 150):
            out_of_spec_count += 1
            
    assert text_count == 1, f"Expected 1 text anomaly, found {text_count}"
    assert out_of_spec_count == 3, f"Expected 3 out-of-spec readings, found {out_of_spec_count}"
    print("XLSX checks passed.")

def verify_docx():
    doc = Document('templates/approval_note.docx')
    text = '\n'.join([p.text for p in doc.paragraphs])
    placeholders = ['{{TITLE}}', '{{EXEC_SUMMARY}}', '{{RECOMMENDATIONS}}', '{{DECISION}}', '{{PREPARED_BY}}']
    for ph in placeholders:
        assert ph in text, f"Missing placeholder {ph} in DOCX"
    print("DOCX checks passed.")

def verify_pptx():
    prs = Presentation('templates/review.pptx')
    assert len(prs.slides) == 7, f"Expected 7 slides, found {len(prs.slides)}"
    print("PPTX checks passed.")

def verify_pdf():
    # Inspection report: true scan (must contain image raster layer, 0 native text to force OCR)
    with open('data/demo_assets/scanned_inspection_report.pdf', 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        assert len(reader.pages) >= 1
        page = reader.pages[0]
        assert len(page.images) > 0, "Scanned inspection report must contain raster image scan layer"
        text = page.extract_text()
        assert len(text.strip()) == 0, "Scanned inspection report should not have selectable native text (forcing OCR cascade)"
    
    # Injected PDF: 7+ pages with injection payload specifically on page 7 (index 6)
    with open('data/demo_assets/injected.pdf', 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        assert len(reader.pages) >= 7, f"Injected PDF must have at least 7 pages, got {len(reader.pages)}"
        text_p7 = reader.pages[6].extract_text()
        assert "SYSTEM COMMAND" in text_p7 or "Ignore prior instructions" in text_p7, "Injection payload not found on page 7"
    print("PDF checks passed.")

def verify_png():
    img1 = Image.open('data/demo_assets/handwritten_note.png')
    assert img1.size == (400, 200)
    
    img2 = Image.open('data/demo_assets/p_and_id_crop.png')
    assert img2.size == (500, 300)
    print("PNG checks passed.")

try:
    verify_xlsx()
    verify_docx()
    verify_pptx()
    verify_pdf()
    verify_png()
    print("All validation checks PASS.")
except AssertionError as e:
    print(f"Validation FAILED: {e}")
except Exception as e:
    print(f"Unexpected error: {e}")
