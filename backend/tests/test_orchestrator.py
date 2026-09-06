"""
Orchestration loop tests.  OWNER: P1.

Each test corresponds to one required behaviour in docs/CONTRACTS.md.  These are
the tests that stop the loop quietly degrading into a hardcoded pipeline.
"""

from __future__ import annotations

import asyncio

import pytest

pytestmark = pytest.mark.filterwarnings("ignore::DeprecationWarning")


# -- helpers -----------------------------------------------------------------


async def _drive(service, text, scenario=None, files=None, approve=True, decisions=None):
    """Create a task, answer every approval, and wait for a terminal state.

    `decisions` optionally supplies a list of booleans, one per approval, so a
    test can approve the plan and then reject a later write.
    """
    from app.contracts import TaskCreateRequest
    from app.orchestration.approvals import ApprovalGate

    record, decision, planner = service.create_task(
        TaskCreateRequest(text=text, scenario=scenario, file_paths=files or [])
    )
    await service.start(record, decision, planner)

    queue = list(decisions or [])
    for _ in range(200):
        await asyncio.sleep(0.01)
        pending = record.pending_approval
        if pending is not None and not pending.decided:
            answer = queue.pop(0) if queue else approve
            ApprovalGate(record).resolve(pending.approval_id, answer)
        if record.is_terminal:
            break
    if record.runner is not None:
        await asyncio.wait_for(record.runner, timeout=5)
    return record


def events_of(service, record, type_):
    return [e for e in service.bus.stream(record.task_id).history() if e.type == type_]


# -- approval gate -----------------------------------------------------------


def test_no_step_runs_before_the_plan_is_approved(service):
    async def scenario():
        from app.contracts import TaskCreateRequest

        record, decision, planner = service.create_task(
            TaskCreateRequest(text="draft an approval note", scenario="flagship")
        )
        await service.start(record, decision, planner)
        for _ in range(100):
            await asyncio.sleep(0.01)
            if record.state == "awaiting_approval":
                break
        assert record.state == "awaiting_approval"
        assert record.observations == []
        assert events_of(service, record, "step_start") == []
        await service.cancel(record)

    asyncio.run(scenario())


def test_rejection_stops_execution_with_no_artifact(service):
    record = asyncio.run(_drive(service, "draft an approval note", "rejection", approve=False))
    assert record.state == "rejected"
    assert record.observations == []
    assert record.artifacts == []


def test_duplicate_approval_cannot_run_a_task_twice(service):
    from app.orchestration.approvals import ApprovalConflict, ApprovalGate

    async def scenario():
        record = await _drive(service, "draft an approval note", "flagship")
        assert record.is_terminal
        steps_before = len(record.observations)
        # Replaying an old approval must be a conflict, not a second run.
        with pytest.raises(ApprovalConflict):
            ApprovalGate(record).resolve("ap_replayed", True)
        await asyncio.sleep(0.05)
        assert len(record.observations) == steps_before

    asyncio.run(scenario())


def test_deciding_the_same_approval_twice_conflicts(service):
    from app.contracts import TaskCreateRequest
    from app.orchestration.approvals import ApprovalConflict, ApprovalGate

    async def scenario():
        record, decision, planner = service.create_task(
            TaskCreateRequest(text="draft an approval note", scenario="flagship")
        )
        await service.start(record, decision, planner)
        for _ in range(100):
            await asyncio.sleep(0.01)
            if record.pending_approval is not None:
                break
        approval_id = record.pending_approval.approval_id
        ApprovalGate(record).resolve(approval_id, True)
        with pytest.raises(ApprovalConflict):
            ApprovalGate(record).resolve(approval_id, True)
        await service.cancel(record)

    asyncio.run(scenario())


def test_a_write_shows_its_actual_content_before_committing(service):
    """Approving a filename is not Layer 4.  The preview carries the content."""

    async def scenario():
        from app.contracts import TaskCreateRequest
        from app.orchestration.approvals import ApprovalGate

        record, decision, planner = service.create_task(
            TaskCreateRequest(text="draft an approval note", scenario="flagship")
        )
        await service.start(record, decision, planner)
        previews = []
        for _ in range(300):
            await asyncio.sleep(0.01)
            pending = record.pending_approval
            if pending is not None and not pending.decided:
                if pending.kind == "write":
                    previews.append(pending.preview or "")
                ApprovalGate(record).resolve(pending.approval_id, True)
            if record.is_terminal:
                break
        await asyncio.wait_for(record.runner, timeout=5)
        assert previews, "docgen must request a write approval"
        assert "holding period of 28 minutes" in previews[0]

    asyncio.run(scenario())


def test_rejecting_a_write_stops_the_task(service):
    record = asyncio.run(
        _drive(service, "draft an approval note", "flagship", decisions=[True, False])
    )
    assert record.state == "rejected"
    assert record.artifacts == []


# -- the flagship, as four visible steps ------------------------------------


def test_flagship_runs_four_steps_and_produces_an_artifact(service):
    record = asyncio.run(_drive(service, "draft an approval note", "flagship"))
    assert record.state == "completed"
    assert [o.target for o in record.observations] == [
        "vision", "kb_search", "reasoning", "docgen",
    ]
    assert record.iterations_used == 4
    assert record.iterations_used <= service.max_iterations
    assert len(record.artifacts) == 1
    assert record.artifacts[0].name.endswith(".docx")


def test_the_checklist_reflects_step_status(service):
    """The UI renders plan.steps, so the executor's status mutations must land on
    the objects the plan holds - not on copies handed out by the planner."""
    record = asyncio.run(_drive(service, "draft an approval note", "flagship"))
    assert [s.status for s in record.plan.steps] == ["done", "done", "done", "done"]
    assert [s.n for s in record.plan.steps] == [1, 2, 3, 4]


def test_a_failed_step_is_marked_failed_not_done(service, monkeypatch):
    from app.mocks import scenarios

    def _failing():
        sc = scenarios.get_scenario("flagship")
        sc.tool_outcomes[("kb_search", 1)] = scenarios.ToolOutcome(
            error_code="TOOL_FAILED", error_message="simulated retrieval failure"
        )
        return sc

    monkeypatch.setitem(scenarios.SCENARIOS, "failing_status", _failing)
    record = asyncio.run(_drive(service, "draft an approval note", "failing_status"))
    statuses = {s.target: s.status for s in record.plan.steps}
    assert statuses["kb_search"] == "failed"
    assert statuses["docgen"] == "done"


def test_the_flagship_docx_is_a_real_openable_file(service):
    from app.tools.base import task_root

    record = asyncio.run(_drive(service, "draft an approval note", "flagship"))
    # ArtifactRef.path is recorded relative to the TASK root, not the shared
    # workspace, so it is joined against the same root the download route
    # resolves against.
    path = task_root(service.settings.workspace, record.task_id) / record.artifacts[0].path
    assert path.is_file() and path.stat().st_size > 5000

    import docx  # python-docx

    document = docx.Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Approval Note" in text
    assert "28 minutes" in text
    assert "SIMULATED" in text.upper()  # mock output is labelled in the document


def test_every_mock_result_is_labelled_simulated(service):
    record = asyncio.run(_drive(service, "draft an approval note", "flagship"))
    agent_obs = [o for o in record.observations if o.kind == "agent"]
    assert agent_obs and all(o.simulated for o in agent_obs)


# -- retries, feedback, escalation ------------------------------------------


def test_low_confidence_retries_with_structured_feedback(service):
    record = asyncio.run(_drive(service, "fix this code", "coding_retry"))
    attempts = events_of(service, record, "attempt")
    assert len(attempts) == 2
    assert attempts[0].data["feedback_injected"] is None
    injected = attempts[1].data["feedback_injected"]
    assert injected and "TypeError" in injected
    assert record.state == "completed"


def test_a_retry_does_not_consume_an_iteration(service):
    record = asyncio.run(_drive(service, "fix this code", "coding_retry"))
    assert record.iterations_used == 1  # two attempts, one dispatched step


def test_exhausted_retries_escalate_to_human_review(service):
    record = asyncio.run(_drive(service, "read this degraded scan", "escalation"))
    assert record.state == "needs_human_review"
    assert record.escalation_reason and "human review" in record.escalation_reason
    attempts = events_of(service, record, "attempt")
    assert len(attempts) == 3  # MAX_ATTEMPTS["vision"]
    assert events_of(service, record, "escalate")


def test_escalation_produces_no_artifact(service):
    record = asyncio.run(_drive(service, "read this degraded scan", "escalation"))
    assert record.artifacts == []


def test_attempt_counting_is_per_step_and_one_indexed(service):
    record = asyncio.run(_drive(service, "read this degraded scan", "escalation"))
    numbers = [e.data["attempt"] for e in events_of(service, record, "attempt")]
    assert numbers == [1, 2, 3]


# -- the loop actually loops -------------------------------------------------


def test_observation_changes_the_next_step(service):
    """The dirty workbook inserts a kb_search that the clean one does not.

    This is a deterministic fixture demonstrating that the executor asks the
    planner for the next action from accumulated observations.  It is not a
    demonstration of live model reasoning.
    """
    dirty = asyncio.run(_drive(service, "check spec", "observation_branch"))
    clean = asyncio.run(_drive(service, "check spec", "observation_branch_clean"))

    dirty_targets = [o.target for o in dirty.observations]
    clean_targets = [o.target for o in clean.observations]

    assert dirty_targets == ["sheet_op", "kb_search", "coding"]
    assert clean_targets == ["sheet_op", "coding"]
    assert dirty.plan.revisions == 1
    assert clean.plan.revisions == 0


def test_a_revision_outside_approved_scope_asks_again(service):
    """kb_search was not in the approved plan for observation_branch, so the
    operator is asked a second time before it runs."""
    record = asyncio.run(_drive(service, "check spec", "observation_branch"))
    requests = events_of(service, record, "approval_request")
    kinds = [r.data["kind"] for r in requests]
    assert "plan" in kinds
    assert "plan_revision" in kinds


def test_rejecting_a_revision_stops_the_task(service):
    record = asyncio.run(
        _drive(service, "check spec", "observation_branch", decisions=[True, False])
    )
    assert record.state == "rejected"
    assert [o.target for o in record.observations] == ["sheet_op"]


# -- bounds ------------------------------------------------------------------


def test_iteration_cap_is_enforced_and_is_not_reported_as_success(service, monkeypatch):
    """A plan longer than the cap ends needs_human_review, never completed."""
    from app.contracts import PlanStep
    from app.mocks import scenarios

    def _long():
        sc = scenarios.get_scenario("flagship")
        sc.initial_steps = [
            PlanStep(n=i, kind="tool", target="list_dir", args={"path": "."},
                     why="step %d" % i)
            for i in range(1, 9)
        ]
        return sc

    monkeypatch.setitem(scenarios.SCENARIOS, "long", _long)
    record = asyncio.run(_drive(service, "long plan", "long"))

    assert record.iterations_used == service.max_iterations
    assert record.state == "needs_human_review"
    assert "iteration cap reached" in (record.escalation_reason or "")


def test_resolve_tool_placeholders_from_prior_observations():
    from app.contracts import Observation, PlanStep
    from app.orchestration.executor import _resolve_tool_placeholders

    step = PlanStep(
        n=2,
        kind="tool",
        target="write_file",
        args={"content": "<generated_code>", "path": "test.py"},
        why="write",
    )
    obs = Observation(
        step_n=1,
        kind="agent",
        target="reasoning",
        ok=True,
        summary="code drafted",
        payload={"content": "def add(a, b):\n    return a + b\n"},
        confidence=0.9,
        iteration=1,
    )
    _resolve_tool_placeholders(step, [obs])
    assert step.args["content"] == "def add(a, b):\n    return a + b\n"


# -- failures become observations, not crashes ------------------------------


def test_a_tool_exception_becomes_an_observation_and_an_event(service, monkeypatch):
    from app.mocks import scenarios

    def _failing():
        sc = scenarios.get_scenario("flagship")
        sc.tool_outcomes[("kb_search", 1)] = scenarios.ToolOutcome(
            error_code="TOOL_FAILED", error_message="simulated retrieval failure"
        )
        return sc

    monkeypatch.setitem(scenarios.SCENARIOS, "failing", _failing)
    record = asyncio.run(_drive(service, "draft an approval note", "failing"))

    failed = [o for o in record.observations if o.target == "kb_search"]
    assert failed and failed[0].ok is False
    assert failed[0].error.code == "TOOL_FAILED"
    assert events_of(service, record, "error")
    # The loop continued past the failure rather than crashing the task.
    assert [o.target for o in record.observations][-1] == "docgen"


# -- audit -------------------------------------------------------------------


def test_every_attempt_is_audited_and_the_chain_verifies(service):
    from app.security.audit import verify_chain

    record = asyncio.run(_drive(service, "fix this code", "coding_retry"))
    entries = service.audit.read_all()
    agent_entries = [e for e in entries if e["action"] == "agent.coding"]
    assert [e["attempt"] for e in agent_entries] == [1, 2]
    assert verify_chain(service.audit.path).ok is True
    _ = record


def test_router_decision_is_audited_before_any_step(service):
    record = asyncio.run(_drive(service, "draft an approval note", "flagship"))
    actions = [e["action"] for e in service.audit.read_all()]
    assert actions[0] == "router.decision"
    assert "plan.proposed" in actions
    _ = record


# -- cancellation ------------------------------------------------------------


def test_cancellation_releases_a_task_waiting_on_approval(service):
    async def scenario():
        from app.contracts import TaskCreateRequest

        record, decision, planner = service.create_task(
            TaskCreateRequest(text="draft an approval note", scenario="flagship")
        )
        await service.start(record, decision, planner)
        for _ in range(100):
            await asyncio.sleep(0.01)
            if record.state == "awaiting_approval":
                break
        assert await service.cancel(record) is True
        assert record.state == "cancelled"
        assert record.runner.done()

    asyncio.run(scenario())


# -- scenario selection ------------------------------------------------------


def test_scenario_selection_uses_whole_words():
    """'inspection' contains 'spec'.  Substring matching sent every
    inspection-report demo to the spreadsheet scenario."""
    from app.mocks.scenarios import select_scenario

    assert select_scenario("Draft an approval note from this inspection report.",
                           "vision", []) == "flagship"
    assert select_scenario("check these readings against spec", "reasoning", []) == \
        "observation_branch"
    assert select_scenario("anything", "reasoning", ["uploads/a.xlsx"]) == \
        "observation_branch"
    assert select_scenario("here is a traceback", "coding", []) == "coding_retry"
    assert select_scenario("this vendor letter looks odd", "reasoning", []) == "injection"
    assert select_scenario("the scan is degraded", "vision", []) == "escalation"
    assert select_scenario("solve this coding problem", "vision", ["problem.png"]) == "vision_code"


def test_every_scenario_key_is_reachable_or_explicitly_manual():
    from app.mocks.scenarios import SCENARIOS

    reachable = {
        "flagship",
        "observation_branch",
        "coding_retry",
        "injection",
        "escalation",
        "vision_code",
    }
    manual = {"rejection", "observation_branch_clean"}  # chosen explicitly in the UI
    assert reachable | manual == set(SCENARIOS)
