import json
import pytest
import asyncio
from unittest.mock import patch, AsyncMock
from pathlib import Path

from app.contracts import SheetOpArgs, ErrorCode, CodingOutput
from app.tools.base import ToolContext, ToolError
from app.config import Settings
from app.tools.sheets import sheet_op
from app.contracts import CodingOutput

@pytest.fixture
def mock_ctx(tmp_path):
    settings = Settings(workspace=tmp_path)
    return ToolContext(settings=settings, task_id="test_task", session_id="test_session")

@pytest.fixture
def sensor_readings(tmp_path):
    import openpyxl
    path = tmp_path / "sensor_readings.xlsx"
    
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Readings"
    ws1.append(["timestamp", "value"])
    ws1.append(["2026-09-01", 12.5])
    ws1.append(["2026-09-02", 9.0])
    ws1.append(["2026-09-03", "error"])
    ws1.append(["2026-09-04", 14.2])
    ws1.append(["2026-09-05", 55.0])
    
    ws2 = wb.create_sheet(title="Spec_Limits")
    ws2.append(["Parameter", "Min", "Max"])
    ws2.append(["Flow Rate", 10.0, 15.0])
    
    wb.save(path)
    return path

@patch("app.tools.sandbox.run_in_sandbox", new_callable=AsyncMock)
def test_compute_sandbox_usage_and_script_generation(mock_run, sensor_readings, mock_ctx):
    # Setup mock to return the expected successful JSON structure
    mock_run.return_value = CodingOutput(
        code="<mock script>",
        stdout=json.dumps({
            "result": {
                "count": 4.0,
                "min": 9.0,
                "max": 55.0,
                "out_of_spec": [
                    {"value": 9.0},
                    {"value": 55.0}
                ]
            },
            "intermediates": {
                "anomalies_found": 1
            }
        }),
        stderr="",
        exit_code=0,
        confidence=1.0,
        sandbox_available=True,
        sandbox_command=[]
    )
    
    spec = json.dumps({
        "action": "analyze",
        "sheet": "Readings",
        "column": "value",
        "min_limit": 10.0,
        "max_limit": 15.0
    })
    
    args = SheetOpArgs(op="compute", path=str(sensor_readings.relative_to(mock_ctx.settings.workspace)), spec=spec)
    
    result = asyncio.run(sheet_op(args, mock_ctx))
    
    # Assert sandbox was called
    assert mock_run.call_count == 1
    call_args = mock_run.call_args[1]
    
    # Verify script generation.  The script must use ONLY what the sandbox
    # image ships (sandbox/Dockerfile pins openpyxl/et-xmlfile/numpy and strips
    # pip): a generated `import pandas` is a guaranteed ModuleNotFoundError.
    generated_code = call_args["code"]
    assert "pandas" not in generated_code
    assert "from openpyxl import load_workbook" in generated_code
    assert "load_workbook" in generated_code
    
    # Verify sandbox usage
    assert result["op"] == "compute"
    assert result["script"] == generated_code
    
    res = result["result"]
    assert res["count"] == 4.0
    assert res["min"] == 9.0
    assert res["max"] == 55.0
    
    oos = res["out_of_spec"]
    assert len(oos) == 2
    
    vals = [row["value"] for row in oos]
    assert 9.0 in vals
    assert 55.0 in vals
    
    assert result["intermediates"]["anomalies_found"] == 1
    
    # Verify source integrity (file wasn't modified or deleted)
    assert sensor_readings.exists()
    
@patch("app.tools.sandbox.run_in_sandbox", new_callable=AsyncMock)
def test_compute_source_is_readonly(mock_run, sensor_readings, mock_ctx):
    mock_run.return_value = CodingOutput(
        code="mock", stdout='{"result": {}}', stderr="", exit_code=0, confidence=1.0, sandbox_available=True, sandbox_command=[]
    )
    
    spec = json.dumps({
        "action": "analyze",
        "sheet": "Readings",
        "column": "value"
    })
    
    args = SheetOpArgs(op="compute", path=str(sensor_readings.relative_to(mock_ctx.settings.workspace)), spec=spec)
    mtime = sensor_readings.stat().st_mtime
    
    asyncio.run(sheet_op(args, mock_ctx))
    
    assert sensor_readings.stat().st_mtime == mtime

def test_compute_invalid_json(sensor_readings, mock_ctx):
    args = SheetOpArgs(op="compute", path=str(sensor_readings.relative_to(mock_ctx.settings.workspace)), spec="not json")
    
    with pytest.raises(ToolError) as excinfo:
        asyncio.run(sheet_op(args, mock_ctx))
        
    assert excinfo.value.code == ErrorCode.INVALID_ARGS
    assert "valid JSON" in excinfo.value.message
    
@patch("app.tools.sandbox.run_in_sandbox", new_callable=AsyncMock)
def test_compute_missing_column(mock_run, sensor_readings, mock_ctx):
    mock_run.return_value = CodingOutput(
        code="mock", stdout='{"error": "Column missing_col not found"}', stderr="", exit_code=0, confidence=1.0, sandbox_available=True, sandbox_command=[]
    )
    
    spec = json.dumps({
        "action": "analyze",
        "sheet": "Readings",
        "column": "missing_col"
    })
    
    args = SheetOpArgs(op="compute", path=str(sensor_readings.relative_to(mock_ctx.settings.workspace)), spec=spec)

    with pytest.raises(ToolError) as excinfo:
        asyncio.run(sheet_op(args, mock_ctx))

    assert excinfo.value.code == ErrorCode.TOOL_FAILED
    assert "not found" in excinfo.value.message


# -- the merged banner row, as shipped in data/demo_assets ---------------------


@pytest.fixture
def banner_workbook(tmp_path):
    """The real demo asset's shape: a merged title row ABOVE the headers."""
    import openpyxl

    path = tmp_path / "banner.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Readings"
    ws.merge_cells("A1:C1")
    ws["A1"] = "UNIT 42 SENSOR READINGS - 2026-09-01"
    ws.append([])  # openpyxl appends after row 1
    ws["A2"], ws["B2"], ws["C2"] = "Timestamp", "Pressure_PSI (PT-101)", "Flow"
    for i in range(5):
        ws.cell(row=3 + i, column=1, value="2026-09-01")
        ws.cell(row=3 + i, column=2, value=155.0 if i == 2 else 100.0)
        ws.cell(row=3 + i, column=3, value=400.0)
    wb.save(path)
    return path


def test_describe_and_read_skip_the_merged_banner_row(banner_workbook, mock_ctx):
    """Row 1 is a merged title; the real headers are on row 2.

    Reading row 1 as the header reports the banner as a column name and shifts
    every value down a row, which is what shipped before this test existed.
    """
    rel = str(banner_workbook.relative_to(mock_ctx.settings.workspace))

    described = sheet_op_sync(SheetOpArgs(op="describe", path=rel), mock_ctx)
    headers = described["schema"]["headers"]["Readings"]
    assert headers[:2] == ["Timestamp", "Pressure_PSI (PT-101)"]
    assert "UNIT 42" not in headers[0]

    got = sheet_op_sync(SheetOpArgs(op="read", path=rel, sheet="Readings"), mock_ctx)
    assert got["headers"][:2] == ["Timestamp", "Pressure_PSI (PT-101)"]
    assert got["rows"][0][1] == 100.0  # a data row, not the header row
    assert got["total_rows"] == 5


def test_generated_script_finds_the_column_under_a_banner_row(banner_workbook):
    """The generated script locates the header row for itself, in-sandbox."""
    from app.tools.sheets import _generate_compute_script

    script = _generate_compute_script(
        json.dumps(
            {
                "action": "analyze",
                "sheet": "Readings",
                "column": "Pressure_PSI (PT-101)",
                "min_limit": 80,
                "max_limit": 150,
            }
        ),
        banner_workbook.name,
    )
    # Run it the way the sandbox would, against the real file.  as_posix()
    # because the substituted path goes into a Python string literal, and the
    # container path this stands in for is always POSIX.
    local = script.replace("/inputs/" + banner_workbook.name, banner_workbook.as_posix())
    namespace: dict = {"__name__": "__main__"}
    import contextlib
    import io

    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        exec(compile(local, "main.py", "exec"), namespace)  # noqa: S102
    payload = json.loads(buf.getvalue())

    assert "error" not in payload
    assert payload["intermediates"]["header_row"] == 2
    assert payload["result"]["count"] == 5.0
    assert [r["_row_index"] for r in payload["result"]["out_of_spec"]] == [5]


def sheet_op_sync(args, ctx):
    return asyncio.run(sheet_op(args, ctx))
