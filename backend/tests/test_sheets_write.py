import pytest
import asyncio
from pathlib import Path

from app.contracts import SheetOpArgs, ErrorCode
from app.tools.base import ToolContext, ToolError
from app.config import Settings
from app.tools.sheets import sheet_op

@pytest.fixture
def mock_ctx(tmp_path):
    settings = Settings(workspace=tmp_path)
    ctx = ToolContext(settings=settings, task_id="test_task", session_id="test_session")
    ctx.ensure_root()
    return ctx

@pytest.fixture
def source_workbook(mock_ctx):
    import openpyxl
    path = mock_ctx.workspace / "sensor_readings.xlsx"
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["id", "value"])
    ws.append([1, 10.0])
    ws.append([2, 20.0])
    
    wb.save(path)
    return path

def test_sheet_write_basic(source_workbook, mock_ctx):
    out_path = "output.xlsx"
    args = SheetOpArgs(
        op="write",
        path=str(source_workbook.relative_to(mock_ctx.workspace)),
        out_path=out_path,
        data={
            "sheet_name": "MyResults",
            "headers": ["ID", "Status"],
            "rows": [[1, "OK"], [2, "OOS"]]
        }
    )
    
    mtime_before = source_workbook.stat().st_mtime
    
    result = asyncio.run(sheet_op(args, mock_ctx))
    
    assert result["op"] == "write"
    assert result["created"] is True
    
    # 2. Output workbook is created
    out_full = mock_ctx.workspace / out_path
    assert out_full.exists()
    
    # 3. Source workbook remains byte-for-byte unchanged (check mtime)
    assert source_workbook.stat().st_mtime == mtime_before
    
    # 7. Output can be reopened with openpyxl
    import openpyxl
    wb = openpyxl.load_workbook(out_full)
    assert "MyResults" in wb.sheetnames
    
    ws = wb["MyResults"]
    assert ws.cell(row=1, column=1).value == "ID"
    assert ws.cell(row=2, column=2).value == "OK"
    assert ws.cell(row=3, column=2).value == "OOS"
    
def test_sheet_write_formatting_status_column(source_workbook, mock_ctx):
    out_path = "output_fmt.xlsx"
    args = SheetOpArgs(
        op="write",
        path=str(source_workbook.relative_to(mock_ctx.workspace)),
        out_path=out_path,
        data={
            "headers": ["Sensor", "Value", "Status"],
            "rows": [
                ["A", 12.5, "OK"],
                ["B", 55.0, "OOS"],
                ["C", 14.2, "OK"]
            ]
        },
        formatting={
            "status_column": "Status",
            "oos_value": "OOS",
            "fill_color": "FFFF0000"
        }
    )
    
    asyncio.run(sheet_op(args, mock_ctx))
    
    out_full = mock_ctx.workspace / out_path
    import openpyxl
    wb = openpyxl.load_workbook(out_full)
    ws = wb.active
    
    # 4. Known OOS values are highlighted
    # Row 1 is headers. Row 2 is A (OK). Row 3 is B (OOS). Row 4 is C (OK).
    # Check row 2 (in-spec)
    assert ws.cell(row=2, column=1).fill.start_color.index != "FFFF0000"
    
    # Check row 3 (OOS) - should be red
    assert ws.cell(row=3, column=1).fill.start_color.index == "FFFF0000"
    assert ws.cell(row=3, column=2).fill.start_color.index == "FFFF0000"
    assert ws.cell(row=3, column=3).fill.start_color.index == "FFFF0000"
    
    # Check row 4 (in-spec)
    assert ws.cell(row=4, column=1).fill.start_color.index != "FFFF0000"
    
def test_sheet_write_formatting_highlight_rows(source_workbook, mock_ctx):
    out_path = "output_fmt2.xlsx"
    args = SheetOpArgs(
        op="write",
        path=str(source_workbook.relative_to(mock_ctx.workspace)),
        out_path=out_path,
        data={
            "headers": ["Sensor", "Value"],
            "rows": [
                ["A", 12.5],
                ["B", 55.0],
                ["C", 14.2]
            ]
        },
        formatting={
            "highlight_rows": [1], # 0-indexed, so row "B"
            "fill_color": "FF00FF00"
        }
    )
    
    asyncio.run(sheet_op(args, mock_ctx))
    
    out_full = mock_ctx.workspace / out_path
    import openpyxl
    wb = openpyxl.load_workbook(out_full)
    ws = wb.active
    
    assert ws.cell(row=2, column=1).fill.start_color.index != "FF00FF00"
    assert ws.cell(row=3, column=1).fill.start_color.index == "FF00FF00" 
    assert ws.cell(row=3, column=1).fill.fill_type == "solid"
    assert ws.cell(row=4, column=1).fill.start_color.index != "FF00FF00"

def test_sheet_write_invalid_color_fails_cleanly(source_workbook, mock_ctx):
    args = SheetOpArgs(
        op="write",
        path=str(source_workbook.relative_to(mock_ctx.workspace)),
        out_path="out.xlsx",
        data={
            "headers": ["A"],
            "rows": [["B"]]
        },
        formatting={
            "fill_color": "invalid_color"
        }
    )
    
    with pytest.raises(ToolError) as excinfo:
        asyncio.run(sheet_op(args, mock_ctx))
        
    assert excinfo.value.code == ErrorCode.INVALID_ARGS
    assert "Invalid fill_color format" in excinfo.value.message
