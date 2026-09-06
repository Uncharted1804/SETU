import os
import hashlib
import shutil
from pathlib import Path
import pytest
from docx import Document

from app.tools.docgen import docgen
from app.contracts import DocgenArgs, ErrorCode
from app.tools.base import ToolContext, ToolError

@pytest.fixture
def docx_template(tmp_path):
    doc = Document()
    doc.add_paragraph('Title: {{TITLE}}')
    doc.add_paragraph('Summary:\n{{EXEC_SUMMARY}}')
    
    table = doc.add_table(rows=1, cols=2)
    cell = table.cell(0, 0)
    cell.text = "Findings:"
    cell = table.cell(0, 1)
    cell.text = "{{FINDINGS_LIST}}"
    
    section = doc.sections[0]
    header = section.header
    header.paragraphs[0].text = "Header: {{HEADER_MARKER}}"
    
    tpl_path = tmp_path / "test_template.docx"
    doc.save(str(tpl_path))
    return tpl_path

@pytest.fixture
def mock_ctx(tmp_path):
    class MockCtx(ToolContext):
        def __init__(self):
            self.mock = False
            self._root = tmp_path
        
        def resolve(self, path: str) -> Path:
            return self._root / path
            
        def new_artifact_name(self, name: str) -> str:
            return "test_generated_" + name
            
        def register_artifact(self, path: Path, media_type: str, simulated: bool):
            class Ref:
                artifact_id = "test_art_id"
                def __init__(self, root, p):
                    self.path = str(p.relative_to(root))
                    self.size_bytes = p.stat().st_size
                    with open(p, "rb") as f:
                        self.sha256 = hashlib.sha256(f.read()).hexdigest()
            return Ref(self._root, path)
            
    return MockCtx()

def test_docgen_docx_full(docx_template, mock_ctx):
    # F. Template integrity - calculate original hash
    with open(docx_template, "rb") as f:
        orig_hash = hashlib.sha256(f.read()).hexdigest()
        
    data = {
        "TITLE": "Test Approval Note",
        "EXEC_SUMMARY": "This is a multiline\nsummary of the inspection.",
        "FINDINGS_LIST": [
            {"text": "Finding 1"},
            "Finding 2"
        ],
        "HEADER_MARKER": "Confidential"
    }
    
    args = DocgenArgs(
        kind="docx",
        data=data,
        template=str(docx_template.name),
        out_name="out.docx"
    )
    
    # Run tool
    res = docgen(args, mock_ctx)
    
    # G. Return contract
    assert res["kind"] == "docx"
    assert res["artifact_id"] == "test_art_id"
    assert "path" in res
    assert res["size_bytes"] > 0
    assert "sha256" in res
    
    # A. Basic generation
    out_path = mock_ctx._root / res["path"]
    assert out_path.exists()
    
    # F. Template integrity check again
    with open(mock_ctx._root / docx_template.name, "rb") as f:
        new_hash = hashlib.sha256(f.read()).hexdigest()
    assert orig_hash == new_hash, "Template was modified!"
    
    # Parse generated doc
    doc = Document(str(out_path))
    
    text = ""
    for p in doc.paragraphs:
        text += p.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text += p.text + "\n"
    for section in doc.sections:
        for p in section.header.paragraphs:
            text += p.text + "\n"
            
    # B. Marker replacement
    assert "{{TITLE}}" not in text
    assert "{{EXEC_SUMMARY}}" not in text
    assert "{{FINDINGS_LIST}}" not in text
    assert "{{HEADER_MARKER}}" not in text
    
    # C. Content correctness
    assert "Test Approval Note" in text
    assert "Confidential" in text
    
    # D. Tables
    assert "Finding 1" in text
    assert "Finding 2" in text
    
    # E. Multiline content
    assert "This is a multiline\nsummary" in text

def test_docgen_invalid_input(mock_ctx):
    # H. Invalid input
    args = DocgenArgs(
        kind="docx",
        data={},
        template="does_not_exist.docx",
        out_name="out.docx"
    )
    with pytest.raises(ToolError) as excinfo:
        docgen(args, mock_ctx)
    assert excinfo.value.code == ErrorCode.NOT_FOUND


def test_docgen_xlsx_with_formatting(mock_ctx):
    from openpyxl import load_workbook
    
    data = {
        "sheet_name": "Sensor_Analysis",
        "headers": ["Timestamp", "Pressure_PSI", "Flow_GPM"],
        "rows": [
            ["2026-09-01 08:00", 100.2, 420.5],
            ["2026-09-01 08:01", 155.0, 422.1],  # Row index 1 (OOS)
            ["2026-09-01 08:02", 102.1, 419.8],
            ["2026-09-01 08:03", 160.2, 425.0],  # Row index 3 (OOS)
        ],
        "formatting": {
            "out_of_spec_rows": [1, 3],
            "fill_color": "FFFFC7CE",
        }
    }
    args = DocgenArgs(
        kind="xlsx",
        data=data,
        out_name="analysis.xlsx"
    )
    res = docgen(args, mock_ctx)
    assert res["kind"] == "xlsx"
    out_path = mock_ctx._root / res["path"]
    assert out_path.exists()
    
    wb = load_workbook(str(out_path))
    ws = wb["Sensor_Analysis"]
    assert ws.max_row == 5  # 1 header + 4 data rows
    
    # Row 2 (index 0) should have no fill
    assert ws.cell(row=2, column=2).fill.fill_type is None
    
    # Row 3 (index 1) was in out_of_spec_rows -> should have solid fill
    assert ws.cell(row=3, column=2).fill.fill_type == "solid"
    assert ws.cell(row=3, column=2).fill.start_color.rgb == "FFFFC7CE"
    
    # Row 4 (index 2) -> normal
    assert ws.cell(row=4, column=2).fill.fill_type is None
    
    # Row 5 (index 3) was in out_of_spec_rows -> solid fill
    assert ws.cell(row=5, column=2).fill.fill_type == "solid"
    assert ws.cell(row=5, column=2).fill.start_color.rgb == "FFFFC7CE"

