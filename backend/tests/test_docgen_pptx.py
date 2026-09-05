import hashlib
from pathlib import Path

import pytest
from pptx import Presentation

from backend.app.contracts import DocgenArgs, ErrorCode
from backend.app.tools.base import ToolContext, ToolError
from backend.app.tools.docgen import docgen

def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

@pytest.fixture
def mock_ctx(tmp_path):
    class MockCtx(ToolContext):
        def __init__(self):
            self.mock = False
            self._root = tmp_path
        
        def resolve(self, path: str) -> Path:
            return self._root / path
            
        def new_artifact_name(self, name: str) -> str:
            return "test_generated_" + name
            
        def register_artifact(self, path: Path, media_type: str, simulated: bool):
            class Ref:
                artifact_id = "test_art_id"
                def __init__(self, root, p):
                    self.path = str(p.relative_to(root))
                    self.size_bytes = p.stat().st_size
                    with open(p, "rb") as f:
                        self.sha256 = hashlib.sha256(f.read()).hexdigest()
            return Ref(self._root, path)
            
    return MockCtx()

def test_pptx_basic_generation_and_contract(mock_ctx, tmp_path):
    # Copy template to tmp_path to avoid modifying the real one even accidentally, 
    # and to ensure it's in the resolve path.
    real_tpl = Path("templates/review.pptx")
    if not real_tpl.is_file():
        pytest.skip("templates/review.pptx is missing")
    
    tpl_path = tmp_path / "review.pptx"
    import shutil
    shutil.copy(real_tpl, tpl_path)
    tpl_hash = _hash(tpl_path)

    args = DocgenArgs(
        kind="pptx",
        template="review.pptx",
        data={
            "PRESENTATION_TITLE": "Automated PPTX Test",
            "DATE": "2026-09-06",
            "EXEC_SUMMARY": "This is a summary.\nIt has multiple paragraphs.",
            "FINDINGS": [
                {"text": "Finding 1", "confidence": 0.9},
                "Finding 2"
            ],
            "DATA_SUMMARY": "Data looks good.",
            "RECOMMENDATIONS": ["Rec 1", "Rec 2"],
            "DECISION": "Approved.",
            "REFERENCES": ["Ref 1", "Ref 2"]
        }
    )

    result = docgen(args, mock_ctx)

    assert result["kind"] == "pptx"
    assert "artifact_id" in result
    assert "size_bytes" in result
    assert "sha256" in result
    
    out_path = mock_ctx.resolve(result["path"])
    assert out_path.is_file()
    assert out_path.name.endswith(".pptx")
    
    # Output separation
    assert out_path.resolve() != tpl_path.resolve()
    
    # Template integrity
    assert _hash(tpl_path) == tpl_hash

    # Valid presentation
    prs = Presentation(str(out_path))
    assert len(prs.slides) > 0

    # Template use / Placeholder replacement
    s0_text = "\n".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "Automated PPTX Test" in s0_text
    assert "2026-09-06" in s0_text
    assert "{{PRESENTATION_TITLE}}" not in s0_text
    
    # Findings/recommendations lists
    s2_text = "\n".join(shape.text for shape in prs.slides[2].shapes if shape.has_text_frame)
    assert "Finding 1" in s2_text
    assert "Finding 2" in s2_text
    assert "- Finding 1" in s2_text
    
def test_pptx_missing_template(mock_ctx):
    args = DocgenArgs(kind="pptx", template="does_not_exist.pptx", data={})
    
    with pytest.raises(ToolError) as exc_info:
        docgen(args, mock_ctx)
        
    assert exc_info.value.code == ErrorCode.NOT_FOUND

def test_pptx_no_template_provided(mock_ctx):
    args = DocgenArgs(kind="pptx", data={})
    
    with pytest.raises(ToolError) as exc_info:
        docgen(args, mock_ctx)
        
    assert exc_info.value.code == ErrorCode.INVALID_ARGS
