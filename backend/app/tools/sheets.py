"""
sheet_op - ONE tool, four operations.  OWNER: P6.  P4 reviews the jail and the
sandbox mount.

Why one tool and not four: the security claim is that the agent's entire
capability surface is SEVEN typed functions.  Four separate spreadsheet tools
would make it ten and the number in the pitch would be wrong.  The operations
are discriminated by a Literal and validated per-operation in
contracts.SheetOpArgs - the public interface is not `**kwargs`.

Why "spreadsheet work" is a TOOL and not an output format: R7 lists it beside
file read/write and sandboxed execution, i.e. among the things the agent calls
MID-TASK.  Producing an .xlsx at the end satisfies R9 and leaves R7 unmet.

    describe  sheet names, dims, headers, dtypes, null counts.  Called FIRST.
              Lets the model reason about a 400-row workbook without loading it
              into an 8k context.  This is the loop earning its keep on stage.
    read      headers + typed rows, truncated.
    compute   runs generated openpyxl IN THE SANDBOX against a READ-ONLY
              mount.  Returns result + script + intermediates - that is where
              "calculations with steps shown" comes from.  The generated script
              uses openpyxl and the standard library ONLY: pandas is not in the
              sandbox image and cannot be installed at runtime.
    write     writes a NEW file into the workspace.  NEVER overwrites an input
              (enforced in SheetOpArgs, not just by convention).

DEPENDENCY NOTE.  `compute` executes inside setu-sandbox:py311, which ships
openpyxl and numpy pinned.  python:3.11-slim does NOT, and nothing is ever
pip-installed during a task.  If the image is missing, compute returns
SANDBOX_UNAVAILABLE - it does not run on the host.
"""

from __future__ import annotations

from typing import Any

from ..contracts import ComputeResult, ErrorCode, SheetOpArgs, SheetRows, SheetSchema, WriteResult
from ..security.paths import PathEscape, to_rel
from .base import ToolContext, ToolError, media_type_for, optional_import


def _open(ctx: ToolContext, rel: str, read_only: bool = True):
    openpyxl = optional_import("openpyxl", owner="P6", purpose="spreadsheet access")
    try:
        target = ctx.resolve(rel)
    except PathEscape as exc:
        raise ToolError(ErrorCode.PATH_ESCAPE, str(exc), requested=rel) from exc
    if not target.is_file():
        raise ToolError(ErrorCode.NOT_FOUND, "no such workbook: " + rel, path=rel)
    return openpyxl.load_workbook(str(target), read_only=read_only, data_only=True), target


def _header_row(ws, max_scan: int = 5) -> int:
    """1-based index of the row holding the column headers.

    A plant export often opens with a merged banner row - one title cell then
    blanks - with the real headers underneath it.  data/demo_assets/
    sensor_readings.xlsx is exactly that shape.  Taking row 1 blindly reports
    the banner as the header and shifts every dtype and value down by a row.
    `read_only=True` does not load merged ranges, so the SHAPE of the row is
    the signal, not the merge itself.
    """
    # ponytail: "first row with >=2 populated cells" heuristic, not a merged-
    # range parser.  Upgrade if a workbook ever ships a real 1-column header.
    for idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1
    ):
        if sum(1 for c in row if c is not None and str(c).strip() != "") >= 2:
            return idx
    return 1


def _describe(args: SheetOpArgs, ctx: ToolContext) -> dict:
    """REAL implementation - describe is on the never-cut list."""
    wb, _ = _open(ctx, args.path)
    sheets: list[str] = []
    dims: dict[str, tuple[int, int]] = {}
    headers: dict[str, list[str]] = {}
    dtypes: dict[str, dict[str, str]] = {}
    nulls: dict[str, dict[str, int]] = {}
    try:
        for ws in wb.worksheets:
            name = ws.title
            sheets.append(name)
            rows = ws.max_row or 0
            cols = ws.max_column or 0
            dims[name] = (rows, cols)
            hrow = _header_row(ws)
            head: list[str] = []
            for row in ws.iter_rows(min_row=hrow, max_row=hrow, values_only=True):
                head = [str(c) if c is not None else "" for c in row]
                break
            headers[name] = head
            col_types: dict[str, str] = {}
            col_nulls: dict[str, int] = {}
            sampled = 0
            for row in ws.iter_rows(
                min_row=hrow + 1, max_row=min(rows, hrow + 200), values_only=True
            ):
                sampled += 1
                for idx, value in enumerate(row):
                    key = head[idx] if idx < len(head) and head[idx] else "col_%d" % (idx + 1)
                    if value is None or value == "":
                        col_nulls[key] = col_nulls.get(key, 0) + 1
                        continue
                    seen = type(value).__name__
                    prior = col_types.get(key)
                    # A column that mixes types is the interesting one - say so.
                    col_types[key] = seen if prior in (None, seen) else "mixed"
            dtypes[name] = col_types
            nulls[name] = col_nulls
    finally:
        wb.close()
    schema = SheetSchema(
        sheets=sheets, dims=dims, headers=headers, dtypes=dtypes, null_counts=nulls
    )
    return {"op": "describe", "path": args.path, "schema": schema.model_dump()}


def _read(args: SheetOpArgs, ctx: ToolContext) -> dict:
    wb, _ = _open(ctx, args.path)
    try:
        ws = wb[args.sheet] if args.sheet else wb.worksheets[0]
        hrow = _header_row(ws)
        head: list[str] = []
        rows: list[list[Any]] = []
        for i, row in enumerate(ws.iter_rows(min_row=hrow, values_only=True), start=hrow):
            if i == hrow:
                head = [str(c) if c is not None else "" for c in row]
                continue
            if len(rows) >= args.max_rows:
                break
            rows.append([c for c in row])
        total = max((ws.max_row or hrow) - hrow, 0)
    finally:
        wb.close()
    out = SheetRows(
        sheet=args.sheet or "sheet0",
        headers=head,
        rows=rows,
        truncated=total > len(rows),
        total_rows=total,
    )
    return {"op": "read", "path": args.path, **out.model_dump()}


def _generate_compute_script(spec_str: str, path_name: str) -> str:
    import json
    try:
        spec = json.loads(spec_str)
    except Exception:
        raise ToolError(ErrorCode.INVALID_ARGS, "spec must be valid JSON")
        
    action = spec.get("action")
    if action != "analyze":
        raise ToolError(ErrorCode.INVALID_ARGS, f"Unsupported compute action: {action}")
        
    sheet = spec.get("sheet", "Readings")
    col = spec.get("column", "value")
    min_limit = spec.get("min_limit")
    max_limit = spec.get("max_limit")
    
    # openpyxl + stdlib ONLY.  setu-sandbox:py311 pins openpyxl, et-xmlfile and
    # numpy (sandbox/Dockerfile) and has pip stripped out, so a generated script
    # that imports pandas fails with ModuleNotFoundError on every real run.
    script = f'''import json
import math
import statistics
from openpyxl import load_workbook

WORKBOOK = "/inputs/{path_name}"
SHEET = {sheet!r}
COLUMN = {col!r}
MIN_LIMIT = {min_limit if min_limit is not None else "None"}
MAX_LIMIT = {max_limit if max_limit is not None else "None"}


def header_row(ws, max_scan=5):
    """First row with two or more populated cells.

    A plant export often opens with a merged banner row - one title cell then
    blanks - with the real headers underneath.  Taking row 1 blindly shifts
    every value down by one row.
    """
    for idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1
    ):
        if sum(1 for c in row if c is not None and str(c).strip() != "") >= 2:
            return idx
    return 1


def cell_value(value):
    """JSON-safe scalar.  Dates and times become ISO strings."""
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value if isinstance(value, int) or math.isfinite(value) else None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def main():
    try:
        wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    except Exception as exc:
        print(json.dumps({{"error": "failed to read workbook: %s" % exc}}))
        return

    if SHEET not in wb.sheetnames:
        print(json.dumps({{
            "error": "sheet %r not found" % SHEET, "sheets": wb.sheetnames
        }}))
        return

    ws = wb[SHEET]
    hrow = header_row(ws)
    headers = []
    for row in ws.iter_rows(min_row=hrow, max_row=hrow, values_only=True):
        headers = [str(c).strip() if c is not None else "" for c in row]
        break

    if COLUMN not in headers:
        print(json.dumps({{
            "error": "Column %s not found" % COLUMN, "headers": headers
        }}))
        return
    ci = headers.index(COLUMN)

    values = []
    non_numeric = 0
    scanned = 0
    out_of_spec = []

    for n, row in enumerate(ws.iter_rows(min_row=hrow + 1, values_only=True), start=hrow + 1):
        if ci >= len(row):
            continue
        scanned += 1
        raw = row[ci]
        if raw is None or str(raw).strip() == "":
            continue
        try:
            value = float(raw)
        except (TypeError, ValueError):
            # The deliberate text value in a numeric column lands here.
            non_numeric += 1
            continue
        if not math.isfinite(value):
            non_numeric += 1
            continue

        values.append(value)
        low = MIN_LIMIT is not None and value < MIN_LIMIT
        high = MAX_LIMIT is not None and value > MAX_LIMIT
        if low or high:
            record = {{}}
            for hi, head in enumerate(headers):
                if head and hi < len(row):
                    record[head] = cell_value(row[hi])
            record["_row_index"] = n
            out_of_spec.append(record)

    result = {{
        "count": float(len(values)),
        "mean": statistics.mean(values) if values else None,
        "min": min(values) if values else None,
        "max": max(values) if values else None,
        "out_of_spec": out_of_spec,
    }}
    intermediates = {{
        "header_row": hrow,
        "rows_scanned": scanned,
        "anomalies_found": non_numeric,
    }}
    print(json.dumps({{"result": result, "intermediates": intermediates}}))


if __name__ == "__main__":
    main()
'''
    return script


async def _compute(args: SheetOpArgs, ctx: ToolContext) -> dict:
    import json
    from pathlib import Path
    from .sandbox import run_in_sandbox

    if not args.spec:
        raise ToolError(ErrorCode.INVALID_ARGS, "compute requires a spec")
        
    path_name = Path(args.path).name
    script = _generate_compute_script(args.spec, path_name)

    out = await run_in_sandbox(
        code=script, 
        settings=ctx.settings, 
        timeout_s=15, 
        input_paths=[args.path]
    )

    if not out.sandbox_available:
        raise ToolError(ErrorCode.SANDBOX_UNAVAILABLE, out.stderr)
        
    if out.exit_code != 0:
        raise ToolError(ErrorCode.TOOL_FAILED, f"Calculation script failed (exit code {out.exit_code}): {out.stderr}")
        
    try:
        parsed = json.loads(out.stdout)
    except json.JSONDecodeError:
        raise ToolError(ErrorCode.TOOL_FAILED, f"Failed to parse calculation output. Raw stdout:\n{out.stdout}")
        
    if "error" in parsed:
        raise ToolError(ErrorCode.TOOL_FAILED, parsed["error"])
        
    res = ComputeResult(
        script=script,
        intermediates=parsed.get("intermediates", {}),
        result=parsed.get("result", {}),
        exit_code=out.exit_code,
        stderr=out.stderr
    )
    
    return {"op": "compute", "path": args.path, **res.model_dump()}


def _write(args: SheetOpArgs, ctx: ToolContext) -> dict:
    """Writes a NEW workbook.  The input is never overwritten."""
    openpyxl = optional_import("openpyxl", owner="P6", purpose="spreadsheet write")
    from openpyxl.styles import PatternFill
    
    try:
        out = ctx.resolve(args.out_path or "")
    except PathEscape as exc:
        raise ToolError(ErrorCode.PATH_ESCAPE, str(exc), requested=args.out_path) from exc
    out.parent.mkdir(parents=True, exist_ok=True)

    data = args.data or {}
    formatting = args.formatting or {}
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(data.get("sheet_name", "Results"))[:31]
    
    headers = data.get("headers")
    if headers:
        ws.append([str(h) for h in headers])
        
    for row in data.get("rows") or []:
        ws.append(list(row))
        
    if formatting:
        color_hex = formatting.get("fill_color", "FFFF0000")
        try:
            fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
        except ValueError:
            raise ToolError(ErrorCode.INVALID_ARGS, f"Invalid fill_color format: {color_hex}")
            
        status_col_name = formatting.get("status_column")
        oos_value = str(formatting.get("oos_value", "OOS"))
        highlight_rows = formatting.get("highlight_rows") or []
        
        status_col_idx = -1
        if headers and status_col_name in headers:
            status_col_idx = headers.index(status_col_name)
            
        start_row = 2 if headers else 1
        for i, row in enumerate(data.get("rows") or []):
            is_oos = False
            if i in highlight_rows:
                is_oos = True
            elif status_col_idx != -1 and status_col_idx < len(row):
                if str(row[status_col_idx]) == oos_value:
                    is_oos = True
                    
            if is_oos:
                for cell in ws[start_row + i]:
                    cell.fill = fill
                    
    wb.save(str(out))

    ref = ctx.register_artifact(out, media_type_for(out), simulated=ctx.mock)
    result = WriteResult(
        path=to_rel(out, ctx.workspace),
        bytes_written=ref.size_bytes,
        created=True,
        sha256=ref.sha256,
    )
    return {"op": "write", "artifact_id": ref.artifact_id, **result.model_dump()}


async def sheet_op(args: SheetOpArgs, ctx: ToolContext) -> dict:
    if args.op == "describe":
        return _describe(args, ctx)
    if args.op == "read":
        return _read(args, ctx)
    if args.op == "compute":
        return await _compute(args, ctx)
    return _write(args, ctx)


_ = ComputeResult  # referenced by the compute contract P6 implements against
