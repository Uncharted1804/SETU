"""
The deliverables engine: docgen(kind, data) -> artifact.  OWNER: P6.

REAL, NOT MOCKED, for docx and xlsx.  The scaffold generates a genuine,
openable .docx from fixture content in mock mode, because a download path that
is never exercised is a download path that breaks at Hour 15.  What is scaffolded
is the POLISH - letterhead, signature block, conditional formatting - not the
file production.

Populating a hand-built template is 30 lines.  Generating a professional-looking
document from scratch at 2 a.m. is not.  P6 builds templates/ at T-3 and this
module then loads them.
"""

from __future__ import annotations

from datetime import datetime, timezone
from typing import Any

from ..contracts import DocgenArgs, ErrorCode
from .base import ToolContext, ToolError, media_type_for, optional_import

DEFAULT_NAMES = {"docx": "approval_note.docx", "xlsx": "calculation.xlsx", "pptx": "review.pptx"}


def _out_path(ctx: ToolContext, args: DocgenArgs):
    name = args.out_name or DEFAULT_NAMES[args.kind]
    if not name.endswith("." + args.kind):
        name = name + "." + args.kind
    rel = "artifacts/" + ctx.new_artifact_name(name)
    target = ctx.resolve(rel)
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def _format_value(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, list):
        lines = []
        for item in val:
            if isinstance(item, dict):
                lines.append("- " + str(item.get("text", item)))
            else:
                lines.append("- " + str(item))
        return "\n".join(lines)
    return str(val)


def _replace_in_paragraph(p, data: dict[str, Any]):
    if not p.text or "{{" not in p.text:
        return
    for k, v in data.items():
        marker = f"{{{{{k}}}}}"
        if marker in p.text:
            val_str = _format_value(v)
            replaced_in_run = False
            for run in p.runs:
                if marker in run.text:
                    run.text = run.text.replace(marker, val_str)
                    replaced_in_run = True
            
            if not replaced_in_run and marker in p.text:
                full_text = p.text.replace(marker, val_str)
                if p.runs:
                    p.runs[0].text = full_text
                    for run in p.runs[1:]:
                        run.text = ""


def _docx(args: DocgenArgs, ctx: ToolContext):
    """A real approval note.  Minimal by design; P6 replaces with the template."""
    docx = optional_import("docx", owner="P6", purpose="DOCX generation")
    data: dict[str, Any] = args.data or {}

    if args.template:
        tpl = ctx.resolve(args.template)
        if not tpl.is_file():
            raise ToolError(ErrorCode.NOT_FOUND, "template not found: " + args.template)
        doc = docx.Document(str(tpl))
        
        for p in doc.paragraphs:
            _replace_in_paragraph(p, data)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _replace_in_paragraph(p, data)
        for section in doc.sections:
            for p in section.header.paragraphs:
                _replace_in_paragraph(p, data)
            for p in section.footer.paragraphs:
                _replace_in_paragraph(p, data)
    else:
        doc = docx.Document()

        doc.add_heading(str(data.get("title", "Approval Note")), level=1)

        meta = data.get("meta") or {}
        if meta:
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for key, value in meta.items():
                row = table.add_row().cells
                row[0].text = str(key)
                row[1].text = str(value)
            doc.add_paragraph("")

        body = data.get("body") or data.get("content") or ""
        for para in str(body).split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())

        findings = data.get("findings") or []
        if findings:
            doc.add_heading("Findings", level=2)
            for f in findings:
                text = f.get("text") if isinstance(f, dict) else str(f)
                conf = f.get("confidence") if isinstance(f, dict) else None
                tier = f.get("extraction_tier") if isinstance(f, dict) else None
                suffix = ""
                if conf is not None:
                    suffix = "  [confidence %.2f%s]" % (conf, ", " + str(tier) if tier else "")
                doc.add_paragraph(str(text) + suffix, style="List Bullet")

        citations = data.get("citations") or []
        if citations:
            doc.add_heading("Sources", level=2)
            for c in citations:
                if isinstance(c, dict):
                    doc.add_paragraph(
                        "%s, p.%s - %s" % (c.get("source_file"), c.get("page"), c.get("snippet", "")[:160]),
                        style="List Number",
                    )
                else:
                    doc.add_paragraph(str(c), style="List Number")

        unsupported = data.get("unsupported_claims") or []
        if unsupported:
            doc.add_heading("Unverified claims - require human confirmation", level=2)
            for claim in unsupported:
                doc.add_paragraph(str(claim), style="List Bullet")

        footer = doc.add_paragraph()
        footer.add_run(
            "Generated by SETU on %s. %s"
            % (
                datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
                "SIMULATED CONTENT - MOCK MODE." if ctx.mock else "Draft for human review.",
            )
        ).italic = True

    target = _out_path(ctx, args)
    doc.save(str(target))
    return target


def _xlsx(args: DocgenArgs, ctx: ToolContext):
    openpyxl = optional_import("openpyxl", owner="P6", purpose="XLSX generation")
    from openpyxl.styles import PatternFill
    
    data: dict[str, Any] = args.data or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(data.get("sheet_name", "Sheet1"))[:31]
    headers = data.get("headers") or []
    if headers:
        ws.append([str(h) for h in headers])
    for row in data.get("rows") or []:
        ws.append(list(row))
        
    formatting = data.get("formatting") or {}
    out_of_spec_rows = formatting.get("out_of_spec_rows") or formatting.get("highlight_rows")
    if out_of_spec_rows:
        color_hex = formatting.get("fill_color", "FFFFC7CE")  # Soft red
        try:
            fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
        except ValueError:
            fill = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
            
        start_row = 2 if headers else 1
        for idx in out_of_spec_rows:
            row_num = idx + start_row if idx < len(data.get("rows") or []) else idx
            if 1 <= row_num <= ws.max_row:
                for cell in ws[row_num]:
                    cell.fill = fill
                    
    target = _out_path(ctx, args)
    wb.save(str(target))
    return target


def _pptx(args: DocgenArgs, ctx: ToolContext):
    pptx = optional_import("pptx", owner="P6", purpose="PPTX generation")
    data: dict[str, Any] = args.data or {}

    if not args.template:
        raise ToolError(ErrorCode.INVALID_ARGS, "pptx generation requires a template")

    tpl = ctx.resolve(args.template)
    if not tpl.is_file():
        raise ToolError(ErrorCode.NOT_FOUND, "template not found: " + args.template)

    prs = pptx.Presentation(str(tpl))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    _replace_in_paragraph(p, data)
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if hasattr(cell, "text_frame") and cell.text_frame:
                            for p in cell.text_frame.paragraphs:
                                _replace_in_paragraph(p, data)

    target = _out_path(ctx, args)
    prs.save(str(target))
    return target


def docgen(args: DocgenArgs, ctx: ToolContext) -> dict:
    builder = {"docx": _docx, "xlsx": _xlsx, "pptx": _pptx}[args.kind]
    target = builder(args, ctx)
    ref = ctx.register_artifact(target, media_type_for(target), simulated=ctx.mock)
    return {
        "kind": args.kind,
        "artifact_id": ref.artifact_id,
        "path": ref.path,
        "size_bytes": ref.size_bytes,
        "sha256": ref.sha256,
    }
